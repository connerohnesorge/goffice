#!/usr/bin/env python3
"""
Create sample .xlsx test files for goffice spreadsheet SDK testing.
"""

from openpyxl import Workbook

def create_minimal_xlsx():
    """Create the simplest valid Excel file with one cell."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws['A1'] = "Hello"
    wb.save('/home/connerohnesorge/Documents/001Repos/goffice/testdata/fixtures/minimal.xlsx')
    print("Created minimal.xlsx")

def create_numbers_xlsx():
    """Create an Excel file with numeric cells and a formula."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    
    # Add numbers 1-5 in column A
    for i in range(1, 6):
        ws.cell(row=i, column=1, value=i)
    
    # Add SUM formula in B1
    ws['B1'] = '=SUM(A1:A5)'
    
    wb.save('/home/connerohnesorge/Documents/001Repos/goffice/testdata/fixtures/numbers.xlsx')
    print("Created numbers.xlsx")

if __name__ == '__main__':
    create_minimal_xlsx()
    create_numbers_xlsx()
